import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch,Mock
from pydantic import ValidationError
from app.document_reader import read_blocks
from app.requirement_catalog import build_catalog,build_catalog_blocks
from app.models import ArchitectureModel,AnalysisResult,FinalModel
from app.validator import validate_architecture,protocols
from app.llm_client import FixtureClient,ReplayClient,LLMClient,LLMError
from app.pipeline import run
from app.evaluation import score,aggregate
from app.renderer import render_mermaid,render_dot

ROOT=Path(__file__).resolve().parents[1]
FIXTURE=ROOT/'evaluation/cases/dev_01.fixture.json'
INPUT=ROOT/'evaluation/cases/dev_01.txt'

def architecture():return ArchitectureModel.model_validate(json.loads(FIXTURE.read_text(encoding='utf-8'))['architecture'])
def sources():return build_catalog(INPUT.read_text(encoding='utf-8'))
def codes(a,entries=None,analysis=None):return {i.code for i in validate_architecture(a,entries or sources(),analysis)}

class CatalogTests(unittest.TestCase):
    def test_wrapped_requirement(self):
        e=build_catalog('3.1 A shall transmit\ndata to B.\n\n3.2 B shall receive.')
        self.assertEqual(len(e),2);self.assertIn('data to B.',e[0].text)
        self.assertEqual(e[0].locations,['line:1','line:2'])
    def test_duplicate_source(self):
        e=build_catalog('3.1 First\n3.1 Second')
        self.assertEqual([x.requirement_id for x in e],['REQ-3.1','REQ-3.1~2'])
        self.assertIn('DUPLICATE_SOURCE_ID',codes(architecture(),e))
    def test_fallback_id_survives_prepend(self):
        a=build_catalog('Alpha\n\nBeta');b=build_catalog('New\n\nAlpha\n\nBeta')
        self.assertEqual(a[1].requirement_id,b[2].requirement_id)
    def test_heading_and_requirement(self):
        e=build_catalog('1. Overview\nG-01: A shall work.\nG-02: B shall work.')
        self.assertEqual(len(e),2);self.assertEqual(e[0].section,'1. Overview')
    def test_docx_order(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'x.docx';doc=Document();doc.add_heading('Section 1',1);doc.add_table(rows=1,cols=1).cell(0,0).text='G-01 | A shall work.';doc.add_heading('Section 2',1);doc.add_paragraph('G-02: B shall work.');doc.save(p)
            b=read_blocks(str(p));self.assertEqual(b[1].kind,'table');self.assertEqual(b[1].section,'Section 1')
            e=build_catalog_blocks(b);self.assertEqual(e[0].locations,['table:1/row:1'])
    def test_blank_pdf_stops(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'blank.pdf';w=PdfWriter();w.add_blank_page(width=100,height=100)
            with p.open('wb') as f:w.write(f)
            with self.assertRaisesRegex(ValueError,'OCR'):read_blocks(str(p))

class ValidatorTests(unittest.TestCase):
    def test_valid(self):self.assertEqual(codes(architecture()),set())
    def test_fake_quote(self):
        a=architecture();a.connections[0].evidence[0].quote='made up'
        self.assertIn('QUOTE_NOT_IN_SOURCE',codes(a))
    def test_wrong_protocol(self):
        a=architecture();a.connections[0].protocol='UDP';self.assertIn('PROTOCOL_MISMATCH',codes(a))
    def test_missing_protocol(self):
        a=architecture();a.connections[0].protocol=None;self.assertIn('EXPLICIT_PROTOCOL_NOT_CAPTURED',codes(a))
    def test_omitted_protocol_in_quote(self):
        a=architecture();a.connections[0].evidence[0].quote='ölçümleri';a.connections[0].protocol=None
        self.assertIn('PROTOCOL_SOURCE_REVIEW',codes(a))
    def test_invalid_coverage(self):
        a=architecture();a.requirement_coverage[0].related_component_ids=['ghost'];a.requirement_coverage[0].requirement_id='GHOST'
        self.assertTrue({'UNKNOWN_COMPONENT_REF','UNKNOWN_COVERAGE_REQUIREMENT'}<=codes(a))
    def test_empty_coverage(self):
        a=architecture();c=a.requirement_coverage[0];c.related_component_ids=[];c.related_connection_ids=[]
        self.assertIn('COVERAGE_WITHOUT_LINKS',codes(a))
    def test_invalid_endpoints(self):
        a=architecture();a.connections[0].target='ghost';self.assertIn('UNKNOWN_TARGET',codes(a))
    def test_duplicates(self):
        a=architecture();a.connections.append(a.connections[0].model_copy(deep=True));self.assertIn('DUPLICATE_CONNECTION_ID',codes(a))
    def test_findings_checked(self):
        a=AnalysisResult(findings=[dict(id='f',severity='high',type='contradiction',title='x',description='x',evidence=[dict(requirement_id='GHOST',quote='x')],related_component_ids=['ghost'])])
        self.assertTrue({'UNKNOWN_COMPONENT_REF','UNKNOWN_EVIDENCE_REF'}<=codes(architecture(),analysis=a))
    def test_missing_quote(self):
        a=architecture();a.components[0].evidence[0].quote=None;self.assertIn('MISSING_QUOTE',codes(a))
    def test_no_substring_can(self):
        self.assertFalse(protocols('An American can operate it.'));self.assertEqual(protocols('CAN bus'),{'CAN'})
    def test_empty_schema_rejected(self):
        with self.assertRaises(ValidationError):ArchitectureModel.model_validate({'system_name':'x'})

class TransportTests(unittest.TestCase):
    def client(self):return LLMClient('https://example.test/v1',api_key='TEST_SECRET',model='fake',max_retries=1)
    def response(self,status=200,content='{"ok":true}',finish='stop'):
        r=Mock();r.status_code=status;r.ok=status<400;r.headers={};r.text='error'
        r.json.return_value={'model':'fake','choices':[{'message':{'content':content},'finish_reason':finish}],'usage':{'total_tokens':4}}
        return r
    def test_json_and_no_secret_in_record(self):
        c=self.client()
        with patch.object(c.session,'post',return_value=self.response()):self.assertEqual(c.extract_json('x','y'),{'ok':True})
        self.assertNotIn('TEST_SECRET',json.dumps(c.calls));self.assertEqual(c.calls[0]['usage']['total_tokens'],4)
    def test_truncation(self):
        c=self.client()
        with patch.object(c.session,'post',return_value=self.response(finish='length')):
            with self.assertRaises(LLMError) as e:c.extract_json('x','y')
        self.assertEqual(e.exception.code,'TRUNCATED')
    def test_429_retry_bound(self):
        c=self.client()
        with patch.object(c.session,'post',return_value=self.response(429)) as p,patch('app.llm_client.time.sleep'):
            with self.assertRaises(LLMError) as e:c.extract_json('x','y')
        self.assertEqual(p.call_count,2);self.assertEqual(e.exception.code,'RATE_LIMIT')
    def test_401_no_retry(self):
        c=self.client()
        with patch.object(c.session,'post',return_value=self.response(401)) as p:
            with self.assertRaises(LLMError):c.extract_json('x','y')
        self.assertEqual(p.call_count,1)
    def test_generic_400_not_retried(self):
        c=self.client();c.json_mode=True
        with patch.object(c.session,'post',return_value=self.response(400)) as p:
            with self.assertRaises(LLMError):c.extract_json('x','y')
        self.assertEqual(p.call_count,1)
    def test_response_format_fallback(self):
        c=self.client();c.json_mode=True;r=self.response(400);r.text='response_format not supported'
        with patch.object(c.session,'post',side_effect=[r,self.response()]) as p:self.assertTrue(c.extract_json('x','y')['ok'])
        self.assertNotIn('response_format',p.call_args.kwargs['json'])
    def test_reasoning_not_final(self):
        with self.assertRaises(LLMError):LLMClient._extract_content({'choices':[{'message':{'reasoning_content':'secret'}}]})
    def test_code_fence(self):self.assertEqual(LLMClient._parse_json('```json\n{"a":1}\n```'),{'a':1})
    def test_prose_rejected(self):
        with self.assertRaises(LLMError):LLMClient._parse_json('Here is {"a":1}')

class PipelineTests(unittest.TestCase):
    def test_offline_and_replay(self):
        with tempfile.TemporaryDirectory() as d:
            one=Path(d)/'one';two=Path(d)/'two'
            a=run(str(INPUT),str(one),FixtureClient(FIXTURE));b=run(str(INPUT),str(two),ReplayClient(one/'run_record.json'))
            self.assertEqual(a.architecture,b.architecture);self.assertEqual(b.review_status,'completed')
            self.assertTrue((one/'source_catalog.json').exists())
    def test_replay_prompt_mismatch(self):
        with tempfile.TemporaryDirectory() as d:
            out=Path(d)/'one';run(str(INPUT),str(out),FixtureClient(FIXTURE));c=ReplayClient(out/'run_record.json')
            with self.assertRaises(LLMError) as e:c.extract_json('changed','prompt')
            self.assertEqual(e.exception.code,'REPLAY_MISMATCH')
    def test_review_failure_preserves_architecture(self):
        client=FixtureClient(FIXTURE);original=client.extract_json
        def extract(s,u):
            if client.index:raise LLMError('RATE_LIMIT','quota')
            return original(s,u)
        client.extract_json=extract
        with tempfile.TemporaryDirectory() as d:
            result=run(str(INPUT),str(Path(d)/'out'),client)
            self.assertEqual(result.review_status,'failed');self.assertTrue(result.architecture.components)
    def test_one_pass_is_skipped(self):
        with tempfile.TemporaryDirectory() as d:self.assertEqual(run(str(INPUT),str(Path(d)/'out'),FixtureClient(FIXTURE),False).review_status,'skipped')
    def test_budget_before_llm(self):
        c=FixtureClient(FIXTURE)
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(LLMError):run(str(INPUT),str(Path(d)/'out'),c,max_input_chars=10)
        self.assertEqual(c.index,0)
    def test_output_reuse_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            Path(d,'old').write_text('old')
            with self.assertRaisesRegex(ValueError,'empty'):run(str(INPUT),d,FixtureClient(FIXTURE))

class ScoringTests(unittest.TestCase):
    def setUp(self):
        self.gold=json.loads((ROOT/'evaluation/cases/dev_01.expected.json').read_text(encoding='utf-8'))
        from dataclasses import asdict
        self.final=FinalModel(architecture=architecture(),analysis=AnalysisResult(),source_catalog=[asdict(e) for e in sources()])
    def test_correct(self):self.assertEqual(score(self.final,self.gold)['connections_with_protocol']['precision'],1)
    def test_reverse(self):
        e=self.final.architecture.connections[0];e.source,e.target=e.target,e.source
        self.assertEqual(score(self.final,self.gold)['directed_connections']['recall'],0)
    def test_wrong_protocol_separate(self):
        self.final.architecture.connections[0].protocol='UDP';s=score(self.final,self.gold)
        self.assertEqual(s['directed_connections']['precision'],1);self.assertEqual(s['connections_with_protocol']['precision'],0)
    def test_duplicate_penalized(self):
        a=self.final.architecture;a.connections.append(a.connections[0].model_copy(deep=True));self.assertEqual(score(self.final,self.gold)['directed_connections']['precision'],.5)
    def test_empty_prediction_not_success(self):
        self.final.architecture.connections=[];s=score(self.final,self.gold)['directed_connections'];self.assertIsNone(s['precision']);self.assertEqual(s['recall'],0)
    def test_failure_visible(self):
        a=aggregate([{'status':'failed','case':'a','repeat':1}]);self.assertEqual(a['run_completion_rate'],0);self.assertIsNone(a['directed_connections']['precision'])
    def test_safe_render_identifiers(self):
        a=self.final.architecture;a.components[0].id='x"] malicious';a.components[0].name='A [x] | "y"';a.connections=[]
        m=render_mermaid(self.final);self.assertNotIn('malicious',m);self.assertIn('#34;',m);self.assertNotIn('malicious',render_dot(self.final))

if __name__=='__main__':unittest.main()
