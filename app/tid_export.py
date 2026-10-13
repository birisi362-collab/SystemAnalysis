"""Export the current normalized TİD as a new document, without rewriting original files."""
from io import BytesIO


def export_tid(project, format='docx', original=False):
    state = project['state']
    entries = project['original']['source_catalog'] if original else state['source_catalog']
    title = project['name']+' — Teknik İsterler Dokümanı'
    subtitle = 'Orijinal belgeden alınan metin' if original else f"Çalışma sürümü {state['document_revision']} · Proje sürümü {project['version']}"
    if format == 'txt':
        lines = [title, subtitle, 'Mühendislik çalışma kopyası', '']
        section = None
        for source in entries:
            if source.get('section') and source['section'] != section:
                section = source['section']; lines.extend([section, ''])
            lines.extend([f"[{source['requirement_id']}] {source['text']}", ''])
        return '\n'.join(lines).encode('utf-8-sig')
    if format != 'docx':
        raise ValueError('DOCX veya TXT biçimini seçin.')
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    doc.sections[0].top_margin = doc.sections[0].bottom_margin = Cm(2)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle)
    doc.add_paragraph('Mühendislik çalışma kopyası. Resmî doküman onayı anlamına gelmez.')
    section = None
    for source in entries:
        if source.get('section') and source['section'] != section:
            section = source['section']; doc.add_heading(section, 1)
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.add_run(source.get('original_id') or source['requirement_id']).bold = True
        doc.add_paragraph(source['text'])
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()
